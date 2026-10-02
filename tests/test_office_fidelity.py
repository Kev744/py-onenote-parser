"""Check semantic fidelity in generated Word and OpenDocument files."""

import struct
import zlib
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document as WordDocument

from pyonenote.models import ContentBlock, Page, Section
from pyonenote.office import write_docx, write_odt

_NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
    "style": "urn:oasis:names:tc:opendocument:xmlns:style:1.0",
    "fo": "urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0",
    "draw": "urn:oasis:names:tc:opendocument:xmlns:drawing:1.0",
    "svg": "urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0",
}


def _section(blocks: list[ContentBlock]) -> Section:
    """Build a section from independently specified content.

    :return: One section with a single page.
    """
    return Section(
        section_id="s",
        title="Section",
        pages=[Page(page_id="p", title="Page", blocks=blocks)],
    )


def _xml(path: Path, member: str) -> ElementTree.Element:
    """Read XML from an Office container.

    :return: The parsed document root.
    """
    with ZipFile(path) as archive:
        return ElementTree.fromstring(archive.read(member))


def test_word_inherited_formatting_and_native_hyperlink(tmp_path: Path) -> None:
    """Retain paragraph defaults, explicit overrides and clickable links."""
    block = ContentBlock(
        text="Plain link",
        style={"bold": True, "font_size": 18, "font": "Arial"},
        children=[
            ContentBlock(kind="run", text="Plain ", style={"bold": False}),
            ContentBlock(
                kind="run", text="link", style={"href": "https://example.org/path"}
            ),
        ],
    )
    destination = tmp_path / "styles.docx"
    write_docx(_section([block]), {}, destination)
    document = WordDocument(str(destination))
    paragraph = document.paragraphs[-1]
    assert paragraph.text == "Plain link"
    assert paragraph.runs[0].bold is False
    font_size = paragraph.runs[0].font.size
    assert font_size is not None and font_size.pt == 18
    assert paragraph.runs[0].font.name == "Arial"
    root = _xml(destination, "word/document.xml")
    link = root.find(".//w:hyperlink", _NS)
    assert link is not None
    assert link.find(".//w:b", _NS) is not None
    rels = _xml(destination, "word/_rels/document.xml.rels")
    assert any(
        item.get("Target") == "https://example.org/path"
        and item.get("TargetMode") == "External"
        for item in rels
    )


def _lists() -> list[ContentBlock]:
    """Build nested-paragraph list items with explicit numbering starts.

    :return: Two ordered lists followed by a bullet list.
    """

    def item(value: str) -> ContentBlock:
        """Wrap a paragraph in a semantic list item."""
        return ContentBlock(kind="list_item", children=[ContentBlock(text=value)])

    return [
        ContentBlock(
            kind="list",
            style={"ordered": True, "start": 5},
            children=[item("Five"), item("Six")],
        ),
        ContentBlock(
            kind="list", style={"ordered": True, "start": 2}, children=[item("Two")]
        ),
        ContentBlock(kind="list", style={"ordered": False}, children=[item("Bullet")]),
    ]


def test_word_lists_number_actual_content_and_restart(tmp_path: Path) -> None:
    """Attach numbering to text paragraphs and use separate restart instances."""
    destination = tmp_path / "lists.docx"
    write_docx(_section(_lists()), {}, destination)
    root = _xml(destination, "word/document.xml")
    numbered = [
        paragraph
        for paragraph in root.findall(".//w:p", _NS)
        if paragraph.find("w:pPr/w:numPr", _NS) is not None
    ]
    assert ["".join(paragraph.itertext()) for paragraph in numbered] == [
        "Five",
        "Six",
        "Two",
        "Bullet",
    ]
    number_ids = []
    for paragraph in numbered:
        number_id = paragraph.find("w:pPr/w:numPr/w:numId", _NS)
        assert number_id is not None
        number_ids.append(number_id.get(f"{{{_NS['w']}}}val"))
    assert number_ids[0] == number_ids[1]
    assert number_ids[2] != number_ids[0]
    numbering = _xml(destination, "word/numbering.xml")
    overrides = numbering.findall(".//w:startOverride", _NS)
    assert {item.get(f"{{{_NS['w']}}}val") for item in overrides} >= {"5", "2"}
    encountered_num = False
    for child in numbering:
        if child.tag == f"{{{_NS['w']}}}num":
            encountered_num = True
        if child.tag == f"{{{_NS['w']}}}abstractNum":
            assert not encountered_num


def test_odt_list_styles_start_and_inherited_runs(tmp_path: Path) -> None:
    """Write explicit ODF list styles and inherited inline typography."""
    block = ContentBlock(
        text="Styled",
        style={"bold": True, "font_size": 17},
        children=[ContentBlock(kind="run", text="Styled")],
    )
    destination = tmp_path / "lists.odt"
    write_odt(_section([block] + _lists()), {}, destination)
    root = _xml(destination, "content.xml")
    assert root.find(".//text:list-level-style-number", _NS) is not None
    assert root.find(".//text:list-level-style-bullet", _NS) is not None
    starts = {
        item.get(f"{{{_NS['text']}}}start-value")
        for item in root.findall(".//text:list-item", _NS)
    }
    assert starts >= {"5", "2"}
    properties = root.findall(".//style:text-properties", _NS)
    assert any(
        item.get(f"{{{_NS['fo']}}}font-weight") == "bold"
        and item.get(f"{{{_NS['fo']}}}font-size") == "17pt"
        for item in properties
    )
    for item in root.findall(".//text:list-item", _NS):
        assert len(item.findall("text:p", _NS)) == 1


def test_office_images_keep_original_aspect_ratio(tmp_path: Path) -> None:
    """Preserve a two-to-one image aspect ratio in both document formats."""

    def chunk(name: bytes, payload: bytes) -> bytes:
        """Build a valid PNG chunk for the independent image fixture."""
        return (
            struct.pack(">I", len(payload))
            + name
            + payload
            + struct.pack(">I", zlib.crc32(name + payload))
        )

    image = tmp_path / "wide.png"
    image.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 1, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00\x00\xff\x00"))
        + chunk(b"IEND", b"")
    )
    section = _section([ContentBlock(kind="image", attachment_id="image")])
    docx = tmp_path / "image.docx"
    odt = tmp_path / "image.odt"
    write_docx(section, {"image": image}, docx)
    write_odt(section, {"image": image}, odt)
    shape = WordDocument(str(docx)).inline_shapes[0]
    assert abs(shape.width / shape.height - 2) < 0.001
    frame = _xml(odt, "content.xml").find(".//draw:frame", _NS)
    assert frame is not None
    width = float(frame.get(f"{{{_NS['svg']}}}width", "0cm").removesuffix("cm"))
    height = float(frame.get(f"{{{_NS['svg']}}}height", "0cm").removesuffix("cm"))
    assert abs(width / height - 2) < 0.001
