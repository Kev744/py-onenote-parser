"""End-to-end real-file regressions for all public export formats."""

import base64
import hashlib
import pickle
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from docx import Document as WordDocument
from odf.opendocument import load as load_odt
from odf import teletype
from odf.text import H, P
from pypdf import PdfReader

from pyonenote import OneNoteExtractor
from pyonenote.models import Document

_FORMATS = ("json", "xml", "pickle", "html", "markdown", "txt", "pdf", "docx", "odt")


def _digest(data: bytes) -> str:
    """Hash original or exported bytes.

    :return: A deterministic SHA-256 digest.
    """
    return hashlib.sha256(data).hexdigest()


def _verify_structure(paths: dict[str, Path], document: Document) -> None:
    """Reopen structural formats and compare hierarchy and attachment payloads.

    :param paths: Export files indexed by suffix.
    """
    restored = Document.model_validate_json(paths[".json"].read_text(encoding="utf-8"))
    assert [
        (section.title, [page.title for page in section.pages])
        for section in restored.sections
    ] == [
        (section.title, [page.title for page in section.pages])
        for section in document.sections
    ]
    assert {
        key: _digest(value.data) for key, value in restored.store.attachments.items()
    } == {key: _digest(value.data) for key, value in document.store.attachments.items()}
    assert restored.store.revisions == document.store.revisions
    for identifier, obj in document.store.objects.items():
        exported = restored.store.objects[identifier]
        assert exported.jcid_name == obj.jcid_name
        assert exported.jcid_aliases == obj.jcid_aliases
        assert exported.jcid_known == obj.jcid_known
        for original, decoded in zip(obj.properties, exported.properties):
            assert decoded.raw == original.raw
            assert decoded.name == original.name
            assert decoded.name_aliases == original.name_aliases
            assert decoded.known == original.known
            assert decoded.interpretation == original.interpretation
            assert decoded.decoded_value == original.decoded_value
            if original.name in {"TopologyCreationTimeStamp", "LastModifiedTimeStamp"}:
                assert decoded.value == original.value
    xml = ElementTree.parse(paths[".xml"]).getroot()
    assert xml.tag == "onenote"
    xml_text = paths[".xml"].read_text(encoding="utf-8")
    assert 'name="jcid_name"' in xml_text
    assert 'name="interpretation"' in xml_text
    for raw_field in xml.findall(".//field[@name='raw']"):
        assert raw_field.text is not None and raw_field.text.startswith("0x")
        bytes.fromhex(raw_field.text[2:])
    for section in document.sections:
        assert section.title in xml_text
        for page in section.pages:
            assert page.title in xml_text
    for attachment in document.store.attachments.values():
        assert base64.urlsafe_b64encode(attachment.data).decode("ascii") in xml_text
    # This pickle was created immediately above by our own exporter.
    assert pickle.loads(paths[".pickle"].read_bytes()) == document


def _verify_content(paths: dict[str, Path], phrase: str) -> PdfReader:
    """Reopen text, PDF and Office documents with their corresponding readers.

    :return: A PDF reader for further media assertions.
    """
    for suffix in (".html", ".md", ".txt"):
        assert phrase in paths[suffix].read_text(encoding="utf-8")
    pdf = PdfReader(paths[".pdf"])
    assert phrase in "\n".join(page.extract_text() for page in pdf.pages)
    word = WordDocument(str(paths[".docx"]))
    assert phrase in "\n".join(paragraph.text for paragraph in word.paragraphs)
    odt = load_odt(str(paths[".odt"]))
    assert phrase in "\n".join(
        teletype.extractText(paragraph)
        for paragraph in odt.getElementsByType(P) + odt.getElementsByType(H)
    )
    return pdf


def _verify_images(
    paths: dict[str, Path], pdf: PdfReader, document: Document, assets: dict[str, str]
) -> None:
    """Verify rendered image resources and unchanged packaged original assets.

    :param document: Parsed source document containing its original media.
    """
    assert any(page.images for page in pdf.pages)
    word = WordDocument(str(paths[".docx"]))
    assert len(word.inline_shapes) >= 1
    original = {
        _digest(document.store.attachments[identifier].data)
        for identifier in assets
        if document.store.attachments[identifier].content_type.startswith("image/")
    }
    with ZipFile(paths[".docx"]) as archive:
        images = {
            _digest(archive.read(name))
            for name in archive.namelist()
            if name.startswith("word/media/")
        }
        assert original == images
    with ZipFile(paths[".odt"]) as archive:
        images = {
            _digest(archive.read(name))
            for name in archive.namelist()
            if name.startswith("Pictures/") and not name.endswith("/")
        }
        assert original == images
    embedded = {_digest(data) for values in pdf.attachments.values() for data in values}
    assert embedded == {
        _digest(document.store.attachments[identifier].data) for identifier in assets
    }


@pytest.mark.parametrize(
    ("filename", "phrase", "has_images"),
    [
        ("Equipe Euro 2016.one", "Cristiano Ronaldo", True),
        ("Section sans titre.one", "history", False),
    ],
)
def test_real_files_export_all_formats_without_source_mutation(
    filename: str, phrase: str, has_images: bool, tmp_path: Path
) -> None:
    """Run the complete public extraction/export flow on both supplied stores."""
    source = Path(__file__).resolve().parents[1] / filename
    original_digest = _digest(source.read_bytes())
    extractor = OneNoteExtractor()
    document = extractor.parse(source)
    result = extractor.export(document, tmp_path, _FORMATS, debug=True)
    paths = {Path(path).suffix: Path(path) for path in result.files}
    assert set(paths) == {
        ".json",
        ".xml",
        ".pickle",
        ".html",
        ".md",
        ".txt",
        ".pdf",
        ".docx",
        ".odt",
    }
    _verify_structure(paths, document)
    pdf = _verify_content(paths, phrase)
    for identifier, asset_path in result.assets.items():
        asset = Path(asset_path)
        assert asset.resolve().is_relative_to(tmp_path.resolve())
        assert _digest(asset.read_bytes()) == _digest(
            document.store.attachments[identifier].data
        )
    if has_images:
        _verify_images(paths, pdf, document, result.assets)
    else:
        assert any(
            item.code == "UNKNOWN_CONTENT_PROPERTY" and "0x240034D8" in item.message
            for item in document.store.diagnostics
        )
    assert _digest(source.read_bytes()) == original_digest
