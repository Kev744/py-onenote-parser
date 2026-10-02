"""Specify readable aggregate text and separate inline/structural content."""

import json
from pathlib import Path
import struct
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document as WordDocument
import pytest

from pyonenote import OneNoteExtractor
from pyonenote.content import assemble_document
from pyonenote.models import ContentBlock, Document, OneObject, ParsedStore, Property
from pyonenote.rendering import block_html, block_text, section_html, section_text


def _object(identifier: str, jcid: int, **values: object) -> OneObject:
    """Build an object using hexadecimal property keys."""
    return OneObject(
        object_id=identifier,
        jcid=jcid,
        properties=[
            Property(property_id=int(key, 16), value=value)
            for key, value in values.items()
        ],
    )


def _document(reverse: bool = False) -> Document:
    """Assemble nested content with independently ordered relationship properties."""
    objects = [
        _object("page", 0x6000B, **{"24001C20": ["outer"]}),
        _object("outer", 0x6000D, **{"24001C1F": ["text"], "24001C20": ["inner"]}),
        _object("inner", 0x6000D, **{"24001C1F": ["nested"]}),
        _object(
            "text",
            0x6000E,
            **{
                "1C001C22": "Hello world",
                "1C001E12": struct.pack("<I", 6),
                "24001E13": ["plain", "bold"],
            },
        ),
        _object("nested", 0x6000E, **{"1C001C22": "Nested content"}),
        _object("plain", 0x12004D),
        _object("bold", 0x12004D, **{"08001C04": True}),
    ]
    if reverse:
        objects[1].properties.reverse()
    return assemble_document(
        ParsedStore(format="revision", objects={obj.object_id: obj for obj in objects})
    )


def test_runs_children_and_live_aggregate_text() -> None:
    """Expose full text at every level without putting inline runs in children."""
    document = _document()
    outer = document.sections[0].pages[0].blocks[0]
    paragraph = outer.children[0]
    assert outer.text == "Hello world\nNested content"
    assert paragraph.text == "Hello world"
    assert paragraph.children == []
    assert [run.text for run in paragraph.runs] == ["Hello ", "world"]
    assert paragraph.runs[1].style == {"bold": True}
    assert paragraph.source_object_id == "text"
    assert "RichEditTextUnicode" not in paragraph.properties
    paragraph.runs[1].text = "reader"
    assert outer.text == "Hello reader\nNested content"
    payload = json.loads(document.model_dump_json())
    assert payload["schema_version"] == 2
    assert (
        Document.model_validate(payload).model_dump(mode="json")["sections"]
        == payload["sections"]
    )


def test_property_order_does_not_change_reading_order() -> None:
    """Read an outline's content before its nested outline elements."""
    assert _document().sections == _document(reverse=True).sections


def test_runs_and_nested_blocks_render_once() -> None:
    """Retain structural descendants in TXT and keep them outside HTML paragraphs."""
    block = ContentBlock.model_validate(
        {
            "kind": "paragraph",
            "runs": [{"text": "Parent"}],
            "children": [{"runs": [{"text": "Child"}]}],
        }
    )
    assert block.text == "Parent\nChild"
    assert block_text(block) == "Parent\nChild"
    html = block_html(block, {}, _document())
    assert html == "<p>Parent</p><p>Child</p>"


def test_legacy_run_children_are_migrated_without_duplicate_text() -> None:
    """Read previous JSON content and never reimport serialized aggregate text."""
    block = ContentBlock.model_validate(
        {
            "text": "ab",
            "children": [
                {"kind": "run", "text": "a"},
                {"kind": "run", "text": "b", "style": {"bold": True}},
            ],
        }
    )
    assert block.children == []
    assert [run.text for run in block.runs] == ["a", "b"]
    assert ContentBlock.model_validate_json(block.model_dump_json()) == block


def test_table_text_retains_empty_cell_positions() -> None:
    """Use tabs between cells and newlines between rows without dropping blanks."""
    row = ContentBlock(
        kind="row",
        children=[ContentBlock(text="A"), ContentBlock(), ContentBlock(text="C")],
    )
    assert row.text == "A\t\tC"
    assert (
        ContentBlock(kind="table", children=[row, row.model_copy(deep=True)]).text
        == "A\t\tC\nA\t\tC"
    )


@pytest.mark.parametrize("filename", ["Equipe Euro 2016.one", "Section sans titre.one"])
def test_real_titles_are_separate_and_body_text_is_not_duplicated(
    filename: str,
) -> None:
    """Preserve title metadata separately from the rendered body."""
    document = OneNoteExtractor().parse(Path(__file__).resolve().parents[1] / filename)
    section = document.sections[0]
    page = section.pages[0]
    assert page.title_block is not None
    assert page.title in page.title_block.text
    assert all(block is not page.title_block for block in page.blocks)
    text = section_text(section)
    assert text.count(page.title) == 1 + int(section.title == page.title)
    assert page.text.count(page.title) == 1


def test_nested_content_exported_once_in_office(tmp_path: Path) -> None:
    """Do not render container summaries in addition to their children."""
    document = _document()
    result = OneNoteExtractor.export(document, tmp_path, ["docx", "odt", "json", "xml"])
    paths = {Path(path).suffix: Path(path) for path in result.files}
    word = WordDocument(str(paths[".docx"]))
    text = "\n".join(paragraph.text for paragraph in word.paragraphs)
    assert text.count("Hello world") == text.count("Nested content") == 1
    with ZipFile(paths[".odt"]) as archive:
        root = ElementTree.fromstring(archive.read("content.xml"))
    text = "".join(root.itertext())
    assert text.count("Hello world") == text.count("Nested content") == 1
    assert ElementTree.parse(paths[".xml"]).getroot().get("schema_version") == "3"


def test_documented_style_units_and_explicit_resets(tmp_path: Path) -> None:
    """Keep paragraph alignment separate from run overrides across exporters."""
    store = _document().store
    spacing = int.from_bytes(struct.pack("<f", 0.5), "little")
    store.objects["default"] = _object(
        "default",
        0x12004D,
        **{
            "14001C0C": 255,
            "08001C04": True,
            "1C001C0A": "Consolas",
            "0C003477": 2,
            "1400342E": spacing,
            "1400342F": spacing,
            "14003430": spacing,
            "08003476": True,
        },
    )
    store.objects["text"].properties.extend(
        [
            Property(property_id=0x2000342C, value="default"),
            Property(property_id=0x0C003477, value=1),
        ]
    )
    store.objects["plain"].properties.extend(
        [
            Property(property_id=0x14001C0C, value=0xFF000000),
            Property(property_id=0x08001C04, value=False),
        ]
    )
    document = assemble_document(store)
    paragraph = document.sections[0].pages[0].blocks[0].children[0]
    assert paragraph.style["alignment"] == "center"
    assert paragraph.style["space_before"] == 18.0
    assert paragraph.style["space_after"] == 18.0
    assert paragraph.style["line_spacing"] == 18.0
    assert paragraph.style["direction"] == "rtl"
    assert paragraph.runs[0].style["color"] == "auto"
    html = block_html(paragraph, {}, document)
    assert "text-align:center" in html and 'dir="rtl"' in html
    assert "font-family:" in html and "Consolas" in html
    assert "color:initial" in html and "margin-top:18.0pt" in html
    assert "<strong>Hello " not in html
    result = OneNoteExtractor.export(document, tmp_path, ["docx", "odt"])
    for path in result.files:
        with ZipFile(path) as archive:
            xml = archive.read(
                "word/document.xml" if path.endswith(".docx") else "content.xml"
            ).decode()
        if path.endswith(".docx"):
            assert 'w:val="center"' in xml and "w:bidi" in xml
            assert 'w:color w:val="auto"' in xml
            assert 'w:before="360"' in xml
        else:
            assert 'fo:text-align="center"' in xml
            assert 'style:writing-mode="rl-tb"' in xml
            assert 'style:use-window-font-color="true"' in xml


def test_unsafe_font_names_cannot_inject_css() -> None:
    """Escape font names as CSS strings and then as HTML attributes."""
    block = ContentBlock(
        text="safe", style={"font": 'A";background:url(https://evil.invalid);/*'}
    )
    html = block_html(block, {}, _document())
    assert 'style="font-family:&quot;' in html
    assert "\\22 " in html


def test_title_formatting_and_metadata_are_not_lost(tmp_path: Path) -> None:
    """Render the title once with its formatting and retain the title date."""
    store = _document().store
    store.objects["page"].properties.append(
        Property(property_id=0x24001D5F, value=["title"])
    )
    store.objects.update(
        {
            "title": _object("title", 0x6002C, **{"24001C20": ["title-text", "date"]}),
            "title-text": _object(
                "title-text",
                0x6000E,
                **{"1C001C22": "Unique title", "08001CB4": True, "08001C05": True},
            ),
            "date": _object(
                "date", 0x6000E, **{"1C001C22": "2026-10-01", "08001CB5": True}
            ),
        }
    )
    document = assemble_document(store)
    section = document.sections[0]
    assert section.pages[0].title == "Unique title"
    html = section_html(section, document, {})
    assert html.count("Unique title") == 1
    assert "<h2><em>Unique title</em></h2>" in html
    assert "2026-10-01" in section.text
    result = OneNoteExtractor.export(document, tmp_path, ["docx", "odt"])
    word = WordDocument(next(path for path in result.files if path.endswith(".docx")))
    paragraph = next(p for p in word.paragraphs if p.text == "Unique title")
    assert paragraph.runs[0].italic
    with ZipFile(
        next(path for path in result.files if path.endswith(".odt"))
    ) as archive:
        xml = archive.read("content.xml").decode()
    assert xml.count("Unique title") == 1 and 'fo:font-style="italic"' in xml


def test_identical_text_at_distinct_locations_is_preserved() -> None:
    """Never remove legitimate repeated paragraphs based on text equality."""
    block = ContentBlock(
        kind="container",
        children=[ContentBlock(text="Again"), ContentBlock(text="Again")],
    )
    assert block.text == "Again\nAgain"
    assert block_html(block, {}, _document()).count("Again") == 2
