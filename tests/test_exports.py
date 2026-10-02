"""Round-trip tests for content, structured exports and attachment safety."""

import json
import pickle
from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from docx import Document as WordDocument
from pypdf import PdfReader

from pyonenote import ReadableDocument
from pyonenote.errors import ExportError
from pyonenote.exports import export_document
from pyonenote.models import (
    Attachment,
    ContentBlock,
    Document,
    Page,
    ParsedStore,
    Section,
)


@pytest.fixture(name="document")
def document_fixture() -> Document:
    """Build known content independently of the binary parsers."""
    return Document(
        store=ParsedStore(
            format="revision",
            attachments={
                "asset": Attachment(
                    attachment_id="asset",
                    file_name="../../CON.txt",
                    data=b"\xff\x00payload",
                )
            },
        ),
        sections=[
            Section(
                section_id="s",
                title="Équipe <2016>",
                pages=[
                    Page(
                        page_id="p",
                        title="Résumé",
                        blocks=[
                            ContentBlock(
                                text="Bonjour été",
                                children=[
                                    ContentBlock(kind="run", text="Bonjour "),
                                    ContentBlock(
                                        kind="run", text="été", style={"bold": True}
                                    ),
                                ],
                            ),
                            ContentBlock(
                                kind="table",
                                children=[
                                    ContentBlock(
                                        kind="row",
                                        children=[
                                            ContentBlock(
                                                kind="cell",
                                                children=[ContentBlock(text="Paris")],
                                            ),
                                            ContentBlock(
                                                kind="cell",
                                                children=[ContentBlock(text="42")],
                                            ),
                                        ],
                                    )
                                ],
                            ),
                            ContentBlock(
                                kind="list",
                                children=[
                                    ContentBlock(kind="list_item", text="Premier")
                                ],
                            ),
                            ContentBlock(
                                kind="attachment",
                                attachment_id="asset",
                                text="Pièce jointe",
                            ),
                        ],
                    )
                ],
            )
        ],
    )


def test_structural_exports_round_trip(document: Document, tmp_path: Path) -> None:
    """Preserve original bytes and hierarchy in all structural formats."""
    result = export_document(document, tmp_path, ["json", "xml", "pickle"])
    json_file = next(Path(path) for path in result.files if path.endswith(".json"))
    parsed = ReadableDocument.model_validate_json(json_file.read_text(encoding="utf-8"))
    assert parsed.sections[0].text == document.sections[0].text
    xml_file = next(Path(path) for path in result.files if path.endswith(".xml"))
    root = ElementTree.parse(xml_file).getroot()
    assert root.tag == "onenote"
    assert "Bonjour été" in xml_file.read_text(encoding="utf-8")
    pickle_file = next(Path(path) for path in result.files if path.endswith(".pickle"))
    assert pickle.loads(pickle_file.read_bytes()) == document
    assert Path(result.assets["asset"]).read_bytes() == b"\xff\x00payload"
    assert Path(result.assets["asset"]).resolve().is_relative_to(tmp_path.resolve())


def test_html_markdown_text_and_office_exports(
    document: Document, tmp_path: Path
) -> None:
    """Reopen exports and check meaningful content, formatting and tables."""
    result = export_document(
        document, tmp_path, ["html", "markdown", "txt", "docx", "odt"]
    )
    paths = {Path(path).suffix: Path(path) for path in result.files}
    html = paths[".html"].read_text(encoding="utf-8")
    assert "Équipe &lt;2016&gt;" in html
    assert "<strong>été</strong>" in html
    assert "<table>" in html and "<li>Premier</li>" in html
    assert "**été**" in paths[".md"].read_text(encoding="utf-8")
    assert "Bonjour été" in paths[".txt"].read_text(encoding="utf-8")
    word = WordDocument(str(paths[".docx"]))
    assert any("Bonjour été" in paragraph.text for paragraph in word.paragraphs)
    assert word.tables[0].cell(0, 0).text == "Paris"
    assert any(run.bold for paragraph in word.paragraphs for run in paragraph.runs)
    with ZipFile(paths[".odt"]) as archive:
        xml = archive.read("content.xml").decode("utf-8")
    assert "Bonjour " in xml and "été" in xml and "table:table" in xml


def test_pdf_contains_text_and_original_attachment(
    document: Document, tmp_path: Path
) -> None:
    """Verify a real rendered PDF, including its embedded binary payload."""
    result = export_document(document, tmp_path, ["pdf"])
    pdf = PdfReader(result.files[0])
    text = "\n".join(page.extract_text() for page in pdf.pages)
    assert "Bonjour été" in text
    assert any(b"\xff\x00payload" in values for values in pdf.attachments.values())


def test_existing_files_are_not_overwritten(document: Document, tmp_path: Path) -> None:
    """Require explicit overwrite even for shared attachment destinations."""
    export_document(document, tmp_path, ["html"])
    with pytest.raises(ExportError, match="exist"):
        export_document(document, tmp_path, ["html"])
    result = export_document(document, tmp_path, ["html"], overwrite=True)
    assert result.files


def test_reject_unknown_format_before_writing(
    document: Document, tmp_path: Path
) -> None:
    """Do not write partial output for an unsupported format name."""
    with pytest.raises(ExportError):
        export_document(document, tmp_path, ["html", "invented"])
    assert not list(tmp_path.iterdir())


def test_diagnostics_remain_in_structural_output(
    document: Document, tmp_path: Path
) -> None:
    """Expose source diagnostics alongside the content being exported."""
    from pyonenote.models import Diagnostic  # pylint: disable=import-outside-toplevel

    document.store.diagnostics.append(
        Diagnostic(code="unsupported", message="Ink retained")
    )
    result = export_document(document, tmp_path, ["json"])
    payload = json.loads(Path(result.files[0]).read_text(encoding="utf-8"))
    assert payload["diagnostics"][0]["code"] == "unsupported"
