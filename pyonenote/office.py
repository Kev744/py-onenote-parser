"""Generate Word and OpenDocument content with their dedicated libraries."""

from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.image.image import Image
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from odf import draw, style, table, text
from odf.opendocument import OpenDocumentText

from pyonenote.models import ContentBlock, Section
from pyonenote.rendering import block_text, page_heading, safe_link


def _word_runs(paragraph: Any, block: ContentBlock) -> None:
    """Add formatted text runs to a Word paragraph."""
    _word_paragraph_style(paragraph, block.style)
    for item in block.runs:
        formatting = block.style | item.style
        run = paragraph.add_run(item.text)
        for name in ("bold", "italic", "underline"):
            setattr(run, name, bool(formatting.get(name, False)))
        run.font.strike = bool(formatting.get("strikethrough", False))
        run.font.superscript = bool(formatting.get("superscript", False))
        run.font.subscript = bool(formatting.get("subscript", False))
        if isinstance(formatting.get("font_size"), (int, float)):
            run.font.size = Pt(max(1, min(200, formatting["font_size"])))
        if formatting.get("font"):
            run.font.name = str(formatting["font"])
        color = str(formatting.get("color", "")).lstrip("#")
        if len(color) == 6 and all(char in "0123456789abcdefABCDEF" for char in color):
            run.font.color.rgb = RGBColor.from_string(color)
        elif color == "auto":
            # No public automatic-color setter is provided by python-docx.
            # pylint: disable=protected-access
            run._r.get_or_add_rPr().get_or_add_color().set(
                qn("w:val"), "auto"
            )  # pylint: disable=protected-access
        link = safe_link(formatting.get("href", ""))
        if link:
            _word_hyperlink(paragraph, run, link)


def _word_paragraph_style(paragraph: Any, formatting: dict[str, Any]) -> None:
    """Apply paragraph alignment, minimum line spacing and reading direction."""
    alignment = {
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }
    if formatting.get("alignment") in alignment:
        paragraph.alignment = alignment[formatting["alignment"]]
    for name in ("space_before", "space_after"):
        if name in formatting:
            setattr(paragraph.paragraph_format, name, Pt(formatting[name]))
    if formatting.get("line_spacing", 0) > 0:
        paragraph.paragraph_format.line_spacing = Pt(formatting["line_spacing"])
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
    if formatting.get("direction") in ("rtl", "ltr"):
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1" if formatting["direction"] == "rtl" else "0")
        # python-docx does not expose paragraph reading direction publicly.
        paragraph._p.get_or_add_pPr().append(bidi)  # pylint: disable=protected-access


def _word_hyperlink(paragraph: Any, run: Any, link: str) -> None:
    """Wrap an existing formatted run in a native external Word hyperlink.

    :param link: A scheme-validated target URL.
    """
    relationship_id = paragraph.part.relate_to(
        link, RELATIONSHIP_TYPE.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    # python-docx exposes hyperlink reading but no public creation method.
    # pylint: disable=protected-access
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)


def _word_numbering(parent: Any, formatting: dict[str, Any]) -> tuple[int, int]:
    """Create a distinct Word numbering instance for one semantic list.

    :return: The numbering identifier and zero-based indentation level.
    """
    root = parent.part.numbering_part.element
    level = max(0, min(8, int(formatting.get("level", 0))))
    abstract_id = (
        max(
            (
                int(item.get(qn("w:abstractNumId")))
                for item in root.findall(qn("w:abstractNum"))
            ),
            default=-1,
        )
        + 1
    )
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    definition = OxmlElement("w:lvl")
    definition.set(qn("w:ilvl"), str(level))
    ordered = formatting.get("ordered") or formatting.get("list_type") == "ordered"
    for tag, value in (
        ("start", "1"),
        ("numFmt", "decimal" if ordered else "bullet"),
        ("lvlText", f"%{level + 1}." if ordered else "•"),
        ("lvlJc", "left"),
    ):
        element = OxmlElement(f"w:{tag}")
        element.set(qn("w:val"), value)
        definition.append(element)
    abstract.append(definition)
    root.insert_element_before(abstract, "w:num")
    number = root.add_num(abstract_id)
    number.add_lvlOverride(ilvl=level).add_startOverride(
        max(1, int(formatting.get("start", 1)))
    )
    return int(number.numId), level


def _word_list(parent: Any, block: ContentBlock, assets: dict[str, Path]) -> None:
    """Apply numbering to actual item content rather than empty wrapper paragraphs."""
    number_id, level = _word_numbering(parent, block.style)
    for child in block.children:
        start = len(parent.paragraphs)
        _word_block(parent, child, assets)
        if len(parent.paragraphs) == start:
            parent.add_paragraph()
        paragraph = parent.paragraphs[start]
        # Direct numbering has no public paragraph API in python-docx.
        # pylint: disable=protected-access
        numbering = (
            paragraph._p.get_or_add_pPr().get_or_add_numPr()
        )  # pylint: disable=protected-access
        numbering.get_or_add_ilvl().val = level
        numbering.get_or_add_numId().val = number_id


def _word_asset(parent: Any, block: ContentBlock, assets: dict[str, Path]) -> None:
    """Embed a proportionally scaled picture or identify its exported attachment."""
    asset = assets.get(block.attachment_id)
    if asset and block.kind == "image":
        parent.add_paragraph().add_run().add_picture(str(asset), width=Inches(5.5))
    elif asset:
        parent.add_paragraph(f"{block.inline_text or asset.name}: assets/{asset.name}")
    else:
        parent.add_paragraph("Attachment unavailable")


def _word_table(parent: Any, block: ContentBlock, assets: dict[str, Path]) -> None:
    """Create a Word table and populate each cell's content."""
    rows = [child for child in block.children if child.kind == "row"]
    columns = max((len(row.children) for row in rows), default=1)
    grid = parent.add_table(rows=len(rows), cols=columns)
    grid.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column, cell in enumerate(row.children):
            target = grid.cell(row_index, column)
            if cell.runs:
                _word_runs(target.paragraphs[0], cell)
            for child in cell.children:
                _word_block(target, child, assets)
            if len(target.paragraphs) > 1 and not target.paragraphs[0].text:
                # python-docx supplies an empty paragraph in every new cell.
                # No public paragraph-removal API is exposed by python-docx.
                # pylint: disable=protected-access
                target._element.remove(target.paragraphs[0]._element)


def _word_block(parent: Any, block: ContentBlock, assets: dict[str, Path]) -> None:
    """Append one recursively structured block to a Word document or cell."""
    if block.kind == "table":
        _word_table(parent, block, assets)
        return
    if block.attachment_id:
        _word_asset(parent, block, assets)
        return
    if block.kind == "list":
        _word_list(parent, block, assets)
        return
    if block.kind in ("container", "row", "cell") or (
        block.kind == "list_item" and not block.runs
    ):
        if block.runs:
            _word_runs(parent.add_paragraph(), block)
        for child in block.children:
            _word_block(parent, child, assets)
        return
    paragraph_style = None
    if block.kind == "list_item":
        paragraph_style = (
            "List Number"
            if block.style.get("list_type") == "ordered" or block.style.get("ordered")
            else "List Bullet"
        )
    if block.kind == "heading":
        paragraph_style = "Heading 3"
    paragraph = parent.add_paragraph(style=paragraph_style)
    _word_runs(paragraph, block)
    for child in block.children:
        _word_block(parent, child, assets)


def write_docx(section: Section, assets: dict[str, Path], destination: Path) -> None:
    """Write section content, tables, formatting and images to a DOCX file."""
    document = WordDocument()
    document.add_heading(section.title, 0)
    for index, page in enumerate(section.pages):
        if index:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _word_runs(document.add_heading(level=1), page_heading(page))
        for block in page.content_blocks:
            _word_block(document, block, assets)
    document.save(str(destination))


def _odt_runs(document: Any, paragraph: Any, block: ContentBlock) -> None:
    """Append ODF spans with corresponding inline text properties."""
    _odt_paragraph_style(document, paragraph, block.style)
    for item in block.runs:
        formatting = block.style | item.style
        attributes = _odt_text_properties(formatting)
        if attributes:
            name = f"T{len(document.automaticstyles.childNodes)}"
            span_style = style.Style(name=name, family="text")
            span_style.addElement(style.TextProperties(**attributes))
            document.automaticstyles.addElement(span_style)
            span = text.Span(stylename=span_style, text=item.text)
        else:
            span = text.Span(text=item.text)
        link = safe_link(formatting.get("href", ""))
        if link:
            anchor = text.A(href=link)
            anchor.addElement(span)
            paragraph.addElement(anchor)
        else:
            paragraph.addElement(span)


def _odt_text_properties(formatting: dict[str, Any]) -> dict[str, str]:
    """Translate merged paragraph/run typography into ODF attributes.

    :return: Text properties retaining explicit false overrides.
    """
    result = {}
    for key, target, enabled, disabled in (
        ("bold", "fontweight", "bold", "normal"),
        ("italic", "fontstyle", "italic", "normal"),
        ("underline", "textunderlinestyle", "solid", "none"),
        ("strikethrough", "textlinethroughstyle", "solid", "none"),
    ):
        if key in formatting:
            result[target] = enabled if formatting[key] else disabled
    for key, target in (
        ("font", "fontfamily"),
        ("color", "color"),
        ("highlight", "backgroundcolor"),
    ):
        if formatting.get(key) == "auto" and key == "color":
            result["usewindowfontcolor"] = "true"
        elif formatting.get(key) == "auto" and key == "highlight":
            result[target] = "transparent"
        elif formatting.get(key):
            result[target] = str(formatting[key])
    if formatting.get("font_size"):
        result["fontsize"] = f'{formatting["font_size"]}pt'
    if formatting.get("superscript") or formatting.get("subscript"):
        result["textposition"] = (
            "super 58%" if formatting.get("superscript") else "sub 58%"
        )
    return result


def _odt_paragraph_style(
    document: Any, paragraph: Any, formatting: dict[str, Any]
) -> None:
    """Apply native ODF paragraph layout without applying it to inline spans."""
    attributes = {}
    if formatting.get("alignment") in ("left", "center", "right"):
        attributes["textalign"] = formatting["alignment"]
    if formatting.get("direction") in ("rtl", "ltr"):
        attributes["writingmode"] = (
            "rl-tb" if formatting["direction"] == "rtl" else "lr-tb"
        )
    for key, target in (
        ("space_before", "margintop"),
        ("space_after", "marginbottom"),
        ("line_spacing", "lineheightatleast"),
    ):
        if key in formatting and (key != "line_spacing" or formatting[key] > 0):
            attributes[target] = f"{formatting[key]}pt"
    if attributes:
        paragraph_style = style.Style(
            name=f"P{len(document.automaticstyles.childNodes)}", family="paragraph"
        )
        paragraph_style.addElement(style.ParagraphProperties(**attributes))
        document.automaticstyles.addElement(paragraph_style)
        paragraph.setAttribute("stylename", paragraph_style)


def _odt_table(document: Any, block: ContentBlock, assets: dict[str, Path]) -> Any:
    """Build an ODF table without flattening its cell hierarchy."""
    grid = table.Table()
    for row in block.children:
        row_element = table.TableRow()
        for cell in row.children:
            cell_element = table.TableCell()
            if cell.runs:
                paragraph = text.P()
                _odt_runs(document, paragraph, cell)
                cell_element.addElement(paragraph)
            for child in cell.children:
                _odt_block(document, cell_element, child, assets)
            row_element.addElement(cell_element)
        grid.addElement(row_element)
    return grid


def _odt_asset(document: Any, block: ContentBlock, assets: dict[str, Path]) -> Any:
    """Embed an image or link to a preserved attachment file."""
    paragraph = text.P()
    asset = assets.get(block.attachment_id)
    if asset and block.kind == "image":
        href = document.addPicture(str(asset))
        image = Image.from_file(str(asset))
        ratio = image.width / image.height
        width = min(14.0, 20.0 * ratio)
        frame = draw.Frame(
            width=f"{width:.6f}cm",
            height=f"{width / ratio:.6f}cm",
            anchortype="as-char",
        )
        frame.addElement(draw.Image(href=href))
        paragraph.addElement(frame)
    elif asset:
        paragraph.addElement(
            text.A(href=f"assets/{asset.name}", text=block.inline_text or asset.name)
        )
    else:
        paragraph.addText("Attachment unavailable")
    return paragraph


def _odt_list(document: Any, block: ContentBlock, assets: dict[str, Path]) -> Any:
    """Build explicitly styled numbered or bulleted lists with restart values.

    :return: An ODF list containing its rendered item paragraphs.
    """
    ordered = block.style.get("ordered") or block.style.get("list_type") == "ordered"
    list_style = text.ListStyle(name=f"L{len(document.automaticstyles.childNodes)}")
    for level in range(1, 10):
        if ordered:
            definition = text.ListLevelStyleNumber(
                level=level, numformat="1", numsuffix=".", startvalue=1
            )
        else:
            definition = text.ListLevelStyleBullet(level=level, bulletchar="•")
        definition.addElement(
            style.ListLevelProperties(
                spacebefore=f"{0.6 * level}cm", minlabelwidth="0.6cm"
            )
        )
        list_style.addElement(definition)
    document.automaticstyles.addElement(list_style)
    listing = text.List(stylename=list_style)
    for index, child in enumerate(block.children):
        attributes = (
            {"startvalue": max(1, int(block.style.get("start", 1)))}
            if ordered and index == 0
            else {}
        )
        item = text.ListItem(**attributes)
        _odt_block(document, item, child, assets)
        listing.addElement(item)
    return listing


def _odt_block(
    document: Any, parent: Any, block: ContentBlock, assets: dict[str, Path]
) -> None:
    """Append one content block to an ODT container."""
    if block.kind == "table":
        parent.addElement(_odt_table(document, block, assets))
        return
    if block.attachment_id:
        parent.addElement(_odt_asset(document, block, assets))
        return
    if block.kind == "list":
        parent.addElement(_odt_list(document, block, assets))
        return
    if block.kind in ("container", "row", "cell") or (
        block.kind == "list_item" and not block.runs
    ):
        if block.runs:
            paragraph = text.P()
            _odt_runs(document, paragraph, block)
            parent.addElement(paragraph)
        for child in block.children:
            _odt_block(document, parent, child, assets)
        return
    paragraph = text.H(outlinelevel=3) if block.kind == "heading" else text.P()
    _odt_runs(document, paragraph, block)
    if block.kind == "unsupported" and not block_text(block):
        paragraph.addText("Unsupported object retained in structural export")
    parent.addElement(paragraph)
    for child in block.children:
        _odt_block(document, parent, child, assets)


def write_odt(section: Section, assets: dict[str, Path], destination: Path) -> None:
    """Write section content and embedded images to an ODT document."""
    document = OpenDocumentText()
    document.text.addElement(text.H(outlinelevel=1, text=section.title))
    for page in section.pages:
        heading = text.H(outlinelevel=2)
        _odt_runs(document, heading, page_heading(page))
        document.text.addElement(heading)
        for block in page.content_blocks:
            _odt_block(document, document.text, block, assets)
    document.save(str(destination), addsuffix=False)
